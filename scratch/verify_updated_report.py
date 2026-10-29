import docx

doc = docx.Document(r'D:\Unified_Healthcare_Project_Report (1).docx')
print("Total Paragraphs:", len(doc.paragraphs))
print("Total Tables:", len(doc.tables))

print("\n--- CHAPTER INTRODUCTIONS CHECK ---")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("CHAPTER "):
        print(f"\n{p.text.strip()}:")
        for j in range(i+1, min(i+5, len(doc.paragraphs))):
            t = doc.paragraphs[j].text.strip()
            if t:
                print(f"  P{j}: {t[:120]}")

print("\n--- TABLES CHECK ---")
for i, table in enumerate(doc.tables):
    print(f"Table {i}: {len(table.rows)} rows, {len(table.columns)} cols. Row 0: {' | '.join(c.text.strip().replace('\n', ' ')[:40] for c in table.rows[0].cells)}")

print("\n--- ALGORITHMS CHECK ---")
for i, p in enumerate(doc.paragraphs):
    if "algorithm " in p.text.lower():
        print(f"P{i}: {p.text.strip()[:100]}")

print("\n--- RECENT UPDATES KEYWORDS CHECK ---")
keywords = ["Fast2SMS", "Dual-Mode", "cancellationReason", "cancellation reason", "FR17", "FR18", "TC20", "TC21", "TC22"]
for kw in keywords:
    found_p = any(kw.lower() in p.text.lower() for p in doc.paragraphs)
    found_t = any(kw.lower() in cell.text.lower() for t in doc.tables for r in t.rows for cell in r.cells)
    print(f"Keyword '{kw}': in paragraphs? {found_p}, in tables? {found_t}")
