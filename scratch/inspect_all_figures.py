import docx

doc = docx.Document(r'D:\Unified_Healthcare_Project_Report (1).docx')

with open('scratch/all_figures_context.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if 'figure ' in txt.lower():
            f.write(f"\n=========================================\n")
            f.write(f"FIGURE PARAGRAPH P{i}: {txt}\n")
            f.write(f"-----------------------------------------\n")
            # print up to 3 paragraphs before
            for b in range(max(0, i-3), i):
                f.write(f"  PREV P{b}: {doc.paragraphs[b].text.strip()[:110]}\n")
            # print up to 3 paragraphs after
            for a in range(i+1, min(len(doc.paragraphs), i+4)):
                f.write(f"  NEXT P{a}: {doc.paragraphs[a].text.strip()[:110]}\n")

print("All figures context dumped.")
