import docx

doc = docx.Document(r'D:\Unified_Healthcare_Project_Report (1).docx')

# Let's map paragraphs and tables in order of document body elements
body = doc.element.body
t_idx = 0
p_idx = 0
elements = []

for child in body:
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        p = doc.paragraphs[p_idx]
        txt = p.text.strip()
        elements.append(('p', p_idx, txt))
        p_idx += 1
    elif tag == 'tbl':
        elements.append(('tbl', t_idx, f"Table {t_idx} ({len(doc.tables[t_idx].rows)} rows)"))
        t_idx += 1

with open('scratch/body_order.txt', 'w', encoding='utf-8') as f:
    for idx, (el_type, num, desc) in enumerate(elements):
        if el_type == 'tbl':
            f.write(f"\n>>> [TABLE {num} AT ELEMENT {idx}]: {desc} <<<\n")
            # print previous 3 paragraphs and next 3 paragraphs
            for back in range(max(0, idx-3), idx):
                if elements[back][0] == 'p':
                    f.write(f"   PREV P{elements[back][1]}: {elements[back][2][:100]}\n")
            for fwd in range(idx+1, min(len(elements), idx+4)):
                if elements[fwd][0] == 'p':
                    f.write(f"   NEXT P{elements[fwd][1]}: {elements[fwd][2][:100]}\n")
            f.write("\n")

print("Body order written.")
