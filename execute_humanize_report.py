import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import shutil
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

target_doc_path = r'D:\Unified_Healthcare_Project_Report (1).docx'
backup_doc_path = r'D:\Unified_Healthcare_Project_Report (1)_Backup.docx'

# 1. Create a backup first
if not os.path.exists(backup_doc_path):
    shutil.copyfile(target_doc_path, backup_doc_path)
    print(f"Created backup at: {backup_doc_path}")

doc = docx.Document(target_doc_path)
print(f"Loaded document with {len(doc.paragraphs)} paragraphs.")

def set_run_arial(run, text, font_size_pt=11.5):
    run.text = text
    run.font.name = 'Arial'
    run.font.size = Pt(font_size_pt)
    r = run._element
    rPr = r.get_or_add_rPr()
    # Ensure w:rFonts is explicitly set for Word rendering
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), 'Arial')
    rFonts.set(qn('w:hAnsi'), 'Arial')
    rFonts.set(qn('w:cs'), 'Arial')

def update_para(p, new_text, font_size_pt=11.5):
    # Check if there is a drawing in this paragraph; do not touch drawings
    if 'w:drawing' in p._element.xml:
        print("Warning: skipping paragraph with drawing")
        return
    
    # Remove all existing runs
    for r in list(p.runs):
        p._element.remove(r._element)
        
    run = p.add_run(new_text)
    run.font.name = 'Arial'
    run.font.size = Pt(font_size_pt)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), 'Arial')
    rFonts.set(qn('w:hAnsi'), 'Arial')
    rFonts.set(qn('w:cs'), 'Arial')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Humanized and Plagiarism-Free Paragraph Map
p_map = {
    # SYNOPSIS
    41: (
        "In modern healthcare environments, patient medical histories remain largely trapped within isolated hospital "
        "and clinic databases. When individuals visit new medical providers or require acute emergency transfers, attending "
        "clinicians are forced to make rapid diagnostic and treatment choices without verified visibility into active medication "
        "regimens, past surgical notes, or critical drug allergies. At the same time, maternal health monitoring and childhood "
        "immunization programs struggle with significant dropout rates due to a reliance on physical tracking cards and the absence "
        "of timely reminders. The Unified Healthcare Appointment and Immunization Tracking System resolves these systemic issues "
        "by establishing a centralized, secure web-based infrastructure that unifies clinical records across medical facilities, "
        "simplifies appointment scheduling, and automates pediatric immunization regimens."
    ),
    42: (
        "The platform coordinates dedicated operational workflows across four primary roles: Patients, Doctors, Hospital Staff, and "
        "System Administrators. Developed using the MERN stack (MongoDB, Express.js, React 18, Node.js) and coupled with a Python Flask "
        "machine learning microservice, the application incorporates verified 14-digit ABHA (Ayushman Bharat Health Account) credentialing, "
        "multilingual user interfaces localized across six Indian regional languages (English, Hindi, Tamil, Telugu, Kannada, and Malayalam), "
        "and dynamic outpatient booking. Incoming appointment requests are evaluated by an ensemble Gradient Boosting triage classifier "
        "to determine clinical urgency (Emergency, Consultation, Follow-up, Vaccination, General Checkup), while real-time workload risk "
        "scoring (0–100) assists clinic managers in balancing daily doctor consultation loads."
    ),
    43: (
        "For urgent inter-hospital transfers, an emergency referral data exchange securely transmits structured diagnostic summaries, "
        "current medications, and known allergies between healthcare facilities in real time. Maternal care is tracked across trimesters "
        "using government-issued RCH / PICME registration identifiers, while an automated pediatric engine calculates developmental vaccine "
        "milestones from birth onward. A multi-channel notification infrastructure combines node-cron background daemons with real-time "
        "Telegram push alerts, mobile SMS, and structured HTML emails to deliver timely booking receipts, 24-hour visit reminders, and dose "
        "schedules. An intelligent cancellation protocol captures patient-provided reasons, updates clinic dashboards immediately, and recycles "
        "the freed consultation slot for waiting patients."
    ),
    44: (
        "The overarching objective of the system is to establish an interoperable, accountable, and accessible digital foundation for "
        "modern clinical administration. By bridging the operational gap between portable electronic medical records, predictive clinical triage, "
        "multi-stakeholder alert gateways, and structured maternal-pediatric care pathways, the proposed architecture substantially reduces "
        "redundant laboratory testing, accelerates emergency handoffs, and improves preventative immunization adherence across participating healthcare centers."
    ),

    # CHAPTER 1: INTRODUCTION
    73: (
        "This chapter examines the practical operational realities, technological gaps, and clinical motivations behind developing a unified "
        "healthcare delivery platform. It highlights how fragmented records, manual scheduling bottlenecks, and missed childhood immunizations "
        "compromise patient outcomes across multi-facility healthcare environments. The discussion outlines the project's primary engineering goals, "
        "user role boundaries, and the full-stack MERN and Python microservice architecture adopted to deliver seamless clinical interoperability."
    ),
    75: (
        "Healthcare delivery in multi-tier hospital systems faces persistent friction caused by decentralized data architectures. Patients "
        "frequently consult medical practitioners across different outpatient clinics, specialized hospitals, and rural health centers "
        "throughout their lives. Because clinical records—including laboratory investigations, medication charts, diagnostic histories, and "
        "allergy profiles—remain locked in disjointed proprietary databases or physical paper folios, attending clinicians must often rely on "
        "unverified verbal recall or repeat expensive diagnostic procedures. During acute emergencies, the absence of instantaneous medical data "
        "access introduces life-threatening treatment delays."
    ),
    76: (
        "Equally pressing is the vulnerability of preventive public health workflows, specifically maternal health and early childhood immunizations. "
        "Although government welfare programs assign unique health identifiers such as the Reproductive and Child Health (RCH) ID and national ABHA "
        "tokens, public health tracking frequently loses momentum between childbirth and subsequent multi-year vaccination milestones. Without proactive, "
        "multi-channel reminder frameworks that reach diverse rural and urban households, immunization dropouts remain common, leaving infants "
        "susceptible to vaccine-preventable illnesses."
    ),
    77: (
        "The Unified Healthcare Appointment and Immunization Tracking System resolves these systemic issues through a centralized, interoperable "
        "web application. It bridges the clinical continuum across four core user roles: patients, consulting doctors, hospital support staff, "
        "and administrative personnel. Patients gain self-service tools to discover available specialists, book confirmed consultation slots, "
        "link their national 14-digit ABHA identity, navigate the portal in their preferred regional Indian language, and monitor childhood vaccination "
        "timelines. Doctors benefit from a unified longitudinal electronic health record (EHR) that aggregates diagnoses and treatments recorded "
        "across all participating network hospitals."
    ),
    78: (
        "To address consultation bottlenecks and operational physician burnout, the platform couples clinical workflows with a dedicated Python "
        "machine learning microservice. Incoming appointment requests undergo automated risk assessment and urgency classification via a Gradient "
        "Boosting model, categorizing visits into Emergency, Consultation, Follow-up, Vaccination, or General Checkup. In parallel, dynamic workload "
        "risk scoring monitors doctor consultation densities in real time, alerting clinic coordinators when scheduled volumes exceed safe clinical limits."
    ),
    79: (
        "Operational continuity is fortified by a multi-channel communication infrastructure. The platform pairs automated node-cron schedule listeners "
        "with instantaneous push alerts dispatched across Telegram bot feeds, mobile SMS, and structured HTML emails. Patients can easily link their "
        "Telegram handle during account setup or profile management to receive real-time schedule updates. When cancellations occur, the platform captures "
        "explicit patient justifications, immediately releases the reserved slot for waiting patients, and alerts clinic personnel, ensuring optimal "
        "hospital resource utilization."
    ),
    81: (
        "Figure 1.1 illustrates the end-to-end clinical workflow, tracing a patient's trajectory from initial registration and ABHA verification "
        "through slot discovery, automated ML triage, physician consultation, and downstream immunization tracking."
    ),
    85: (
        "Specific engineering objectives include: (1) Constructing an intelligent cancellation engine that captures structured patient reasoning, "
        "immediately recycles open appointment slots, and updates clinical staff; (2) Establishing a multi-channel notification pipeline combining "
        "real-time Telegram bot push alerts, mobile SMS, and responsive HTML emails for instantaneous booking confirmations and scheduled reminders; "
        "and (3) Integrating government RCH tracking with pediatric vaccination milestone calculators to safeguard child health."
    ),
    87: (
        "The primary goal of this project is to build an interconnected digital healthcare ecosystem that eliminates data fragmentation across "
        "medical institutions. By providing a secure, role-governed electronic health record (EHR) accessible to authenticated medical professionals, "
        "the platform ensures that critical clinical history—including active prescriptions, prior diagnoses, and documented allergies—travels "
        "seamlessly alongside the patient, significantly accelerating emergency interventions and eliminating duplicate laboratory tests."
    ),
    88: (
        "A secondary core objective is institutionalizing proactive preventive healthcare for mothers and infants. By integrating government health "
        "identifiers (RCH / PICME and ABHA IDs) directly with programmatic dosage schedules and background notification daemons, the system ensures "
        "consistent antenatal monitoring and pediatric vaccination adherence. Furthermore, interface localization across six Indian languages removes "
        "linguistic barriers for diverse patient populations, while machine learning triage models prevent clinical schedule congestion by balancing "
        "doctor consultation workloads."
    ),
    90: (
        "The engineering scope covers the design and deployment of a full-stack web architecture utilizing MongoDB, Express.js, React 18, and "
        "Node.js, tailored to support four operational user roles: Patients, Clinicians, Support Staff, and Hospital Administrators. Key functional "
        "capabilities include centralized electronic health record management, cross-hospital emergency referral data dispatching, national ABHA "
        "digital health credentialing, and secure role-based session authorization."
    ),
    92: (
        "Within preventive healthcare, the platform encompasses maternal health registries linked to government RCH references and an automated "
        "pediatric immunization engine that calculates developmental vaccination milestones. Multi-channel alert gateways dispatch automated "
        "reminders via background cron tasks, direct Telegram bot notifications, SMS, and transactional emails. The clinical intelligence layer "
        "incorporates an asynchronous Python Flask microservice executing Gradient Boosting Classifiers for symptom triage and doctor workload risk "
        "scoring. Comprehensive security practices, including bcrypt password salting, JSON Web Token validation, and encrypted payload transmission, "
        "govern all system transactions."
    ),

    # CHAPTER 2: SYSTEM ANALYSIS
    99: (
        "This chapter assesses existing operational workflows within traditional outpatient and hospital management systems. It identifies critical "
        "vulnerabilities surrounding record fragmentation, manual scheduling bottlenecks, untracked childhood immunizations, and communication gaps "
        "between healthcare centers. Drawing upon this analysis, the section introduces the proposed unified platform, highlighting its operational "
        "advantages and detailing the formal functional and non-functional requirements necessary for reliable clinical deployment."
    ),
    101: (
        "In conventional healthcare settings, patient medical histories are typically maintained within isolated electronic health record software "
        "or physical paper charts restricted to a single hospital facility. When a patient switches medical centers or is transported during an acute "
        "emergency, attending doctors have no direct channel to inspect historical diagnoses, current drug prescriptions, previous surgical notes, or "
        "critical drug allergies. The lack of standardized digital health identifier integration further prevents automated record matching across institutions."
    ),
    102: (
        "Similarly, maternal monitoring and child vaccination programs depend heavily on physical immunization cards or local manual ledgers. Due to the "
        "absence of automated multi-channel reminder alerts, mothers frequently miss scheduled antenatal checkups and childhood vaccine milestones. In "
        "addition, rural patient interactions are impaired by digital interfaces that lack regional language localization. Standard clinic appointment "
        "scheduling relies primarily on unmanaged first-come-first-served queues, failing to account for clinical urgency or physician workload saturation."
    ),
    103: (
        "Consequently, traditional healthcare environments operate as isolated silos, where record portability, emergency inter-hospital referrals, "
        "preventive immunization tracking, and clinical workload management function without coordinated technological support."
    ),
    105: "Isolated patient health records: Medical histories remain locked within single-facility databases, creating severe clinical diagnostic blind spots.",
    106: "Absence of digital health identification: Missing ABHA linkage prevents automated patient identity resolution across healthcare centers.",
    107: "Delayed emergency handoffs: Transferring critical care patients occurs without immediate electronic transmission of vital signs and allergy histories.",
    108: "Missed childhood vaccinations: Reliance on physical cards without automated cron and Telegram notifications causes high immunization default rates.",
    109: "Redundant clinical diagnostics: Clinicians frequently repeat blood analyses and radiologic scans due to unavailable historical laboratory data.",
    110: "Linguistic accessibility barriers: Web portals without regional Indian language localization exclude digitally marginalized patient groups.",
    111: "Unstratified appointment queues: Consultation slots are booked without preliminary algorithmic evaluation of clinical symptom severity.",
    112: "Disproportionate doctor workload: Consultation schedules fail to account for cumulative case complexity, contributing to clinician fatigue.",
    113: (
        "To overcome these operational barriers, the platform deploys a multi-channel alert gateway managed by background scheduler workers and "
        "dedicated notification services. The framework dispatches real-time booking confirmations, 24-hour reminder alerts, and cancellation notices "
        "across three distinct channels: real-time Telegram bot push messages, mobile SMS, and structured HTML emails. Patients can link their Telegram "
        "account during onboarding or update preferences via their profile, ensuring dependable message delivery. Furthermore, when an appointment is "
        "cancelled, the application requires a clear cancellation justification (via pre-configured reason chips or custom input), immediately recycles "
        "the open slot back into the doctor's available pool, and dispatches real-time notifications to the affected physician and clinic staff."
    ),
    115: (
        "The proposed Unified Healthcare Appointment and Immunization Tracking System establishes an integrated, web-based digital platform engineered "
        "to overcome legacy healthcare fragmentation. By coupling an event-driven MERN stack application with an asynchronous Python Flask machine "
        "learning microservice, the platform consolidates patient electronic health records, validates national ABHA credentials, manages outpatient "
        "booking schedules, tracks maternal-child immunization regimens, and employs machine learning for automated triage and doctor workload optimization."
    ),
    116: (
        "The system architecture cleanly demarcates responsibilities across four distinct layers: a responsive React 18 presentation client supporting "
        "6-language i18n localization; an Express.js REST API layer governing authentication, session routing, and background notification schedulers; a "
        "schema-governed MongoDB document persistence layer; and an isolated Python Flask microservice executing trained Gradient Boosting inference "
        "models. This modular decoupling ensures robust data persistence, high concurrency handling during peak morning booking hours, and sub-second "
        "machine learning predictions."
    ),
    118: "User onboarding and role-governed authorization across Patient, Doctor, Support Staff, and Hospital Administrator accounts.",
    119: "Digital ABHA (Ayushman Bharat Health Account) verification and electronic credential binding using sandbox specifications.",
    120: "Dynamic multilingual user interface supporting instantaneous localization across six regional languages.",
    121: "Specialist discovery, real-time consultation slot viewing, and self-service online appointment reservations.",
    122: "Longitudinal electronic health record consolidation aggregating past diagnoses, prescriptions, and laboratory reports.",
    123: "Inter-hospital emergency referral data sharing for transferring clinical summaries and allergy profiles during acute care handoffs.",
    124: "Maternal health monitoring linked with government-issued RCH / PICME registration numbers.",
    125: "Pediatric immunization tracking with automated milestone calculation and multi-channel Telegram/SMS/email reminders.",
    126: "Machine learning symptom triage using Gradient Boosting Classifiers to categorize appointment urgency levels.",
    127: "Physician workload risk score evaluation (0–100) to optimize clinical consultation distribution and prevent burnout.",
    129: "Establishes a unified, interoperable platform consolidating clinical records, digital identity, and preventive care workflows.",
    130: "Eliminates departmental record silos by maintaining continuous longitudinal health histories accessible to authorized clinicians.",
    131: "Incorporates national ABHA identifiers to ensure standardized citizen identity resolution across healthcare institutions.",
    132: "Accelerates acute clinical decision-making during emergency transfers by providing instant access to vital summaries and drug allergies.",
    133: "Improves pediatric and maternal health outcomes through automated milestone calculations and proactive multi-channel notification alerts.",
    134: "Overcomes digital divides through comprehensive regional language localization spanning six Indian languages.",
    135: "Applies machine learning to stratify appointment urgency, ensuring high-risk patients receive prioritized medical attention.",
    136: "Protects healthcare professionals from clinical fatigue by providing algorithmic workload distribution analytics.",
    137: "Substantially lowers healthcare costs and patient inconvenience by eliminating unnecessary duplicate diagnostic tests.",
    138: "Follows modular RESTful microservice engineering standards, separating user interface rendering from heavy computational logic.",
    139: "Establishes a robust architectural framework ready for future expansion into IoT biometric telemetry and automated radiology screening.",

    # CHAPTER 3: SYSTEM DESIGN
    150: (
        "This chapter details the architectural blueprints, structural design paradigms, and data models that underpin the unified healthcare platform. "
        "It provides comprehensive multi-tier architectural schematics, operational workflow sequences, and formal UML Activity and BPMN process "
        "representations. Additionally, it specifies the MongoDB document schema configurations, entity relationship definitions, and the integrated "
        "machine learning pipeline governing clinical triage."
    ),
    152: (
        "The Unified Healthcare Appointment and Immunization Tracking System is organized around a multi-tier microservice architecture. The React 18 "
        "frontend acts as the client presentation layer, styled with TailwindCSS and localized using react-i18next. The Node.js and Express.js backend "
        "orchestrates core business logic, session validation, ABHA Sandbox integration, and background notification schedulers. An independent Python "
        "Flask service performs machine learning inference, while MongoDB Atlas serves as the unified document persistence layer."
    ),
    158: "Presentation Layer – React 18 single-page application engineered with TailwindCSS, delivering accessible UI workflows and instant 6-language translation.",
    159: "Application / API Layer – Node.js and Express.js REST web services managing authentication middleware, notification schedulers, and database interactions.",
    160: "Machine Learning Microservice Layer – Python Flask service executing pre-trained Gradient Boosting models for clinical urgency triage and workload risk estimation.",
    161: "Data Persistence Layer – MongoDB database maintaining flexible document collections for patient EHRs, doctor slots, appointments, and immunization batches.",
    162: "Integration & Security Layer – National ABHA sandbox integration, JWT session management, bcryptjs password hashing, and role-based route guards.",
    163: "Visualization Layer – Clinical dashboard widgets, doctor workload risk meters, immunization milestone timelines, and appointment audit tooltips.",
    164: "Architecture Flow: React 18 UI (Vite / i18n) → Express REST API (Node.js / JWT) → Mongoose Validation → MongoDB Atlas Persistence / Flask ML Microservice (Gradient Boosting) → Real-Time Analytics & Telegram/Email Notification Dispatch",
    165: "Algorithm 2 defines the sequential procedure for validating cancellation time windows, updating slot availability, and triggering multi-channel notifications.",
    168: (
        "The operational system workflow encompasses user registration, digital ABHA credential verification, regional language selection, doctor slot "
        "generation, online appointment booking, machine learning urgency triage, clinical consultation recording, maternal tracking under government RCH "
        "identifiers, pediatric vaccination milestone logging, and automated background reminder dispatching."
    ),
    170: "Patient initiates an appointment booking request by selecting a doctor, open slot, and submitting clinical symptoms.",
    171: "Express backend extracts clinical parameters and transmits an asynchronous inference payload to the Flask ML microservice.",
    172: "Flask microservice applies pre-trained LabelEncoders and executes a Gradient Boosting Classifier for urgency category prediction.",
    173: "ML service computes an individualized clinical risk score based on patient age, conditions, and symptom severity.",
    174: "Appointment is recorded in MongoDB with predicted urgency classification and risk attributes, triggering confirmation alerts.",
    175: "Doctor consultation dashboard retrieves updated patient queues and recalculates the physician's cumulative workload risk index.",
    176: "During clinical encounters, the attending doctor inspects the patient's consolidated longitudinal health history across all participating facilities.",
    177: "Doctor documents new clinical findings, active diagnoses, and digital prescriptions into the unified electronic health record.",
    178: "For maternal patients, consultations and antenatal vitals are linked under their official government RCH / PICME identifier.",
    179: "Pediatric patient registration automatically generates developmental vaccination milestone dates and schedules.",
    180: "Automated background schedulers (node-cron) continuously evaluate upcoming due dates to dispatch Telegram, SMS, and email alerts.",
    182: "Patient logs in, authenticates via JWT, selects a consulting specialist, chooses an available time slot, and confirms the visit.",
    183: "Consulting doctor logs in, reviews the daily schedule, inspects historical medical records, and enters consultation outcomes.",
    184: "Hospital administrator logs in to manage doctor credentials, configure department consultation fees, and generate booking slots.",
    185: "Backend API routes process incoming payloads, validate data structures via Mongoose schemas, and update MongoDB document stores.",
    186: "Client dashboards refresh dynamically through reactive state management, rendering updated metrics and status indicators.",
    191: "Figure 3.2 illustrates the UML activity diagram modeling patient decision flows from login and doctor discovery to appointment booking and cancellation.",
    193: (
        "The UML activity workflow models user authentication, session role identification, and primary operational paths. Patients discover physicians, "
        "reserve open slots, inspect consolidated medical records, and track vaccination due dates. Clinicians inspect daily consultation queues, review "
        "historical treatments, record diagnoses, and evaluate clinical workload indices. Administrators onboard practitioners and manage slot templates, "
        "with all operations updating MongoDB and triggering appropriate UI notifications."
    ),
    201: "Figure 3.3 presents the BPMN process model detailing cross-departmental coordination among patients, doctors, hospital staff, and system administrators.",
    203: "BPMN Process Explanation:",
    204: "Step 1 – Registration & Authentication: User creates credentials, verifies ABHA ID, and receives a signed JWT authentication token.",
    205: "Step 2 – Appointment Booking: Patient identifies consulting doctor, picks an open slot, and provides clinical symptom details.",
    206: "Step 3 – ML Triage Execution: Flask microservice executes Gradient Boosting inference to predict urgency category and risk score.",
    207: "Step 4 – Workload Score Calculation: Assigned physician's cumulative workload risk index is recalculated in real time.",
    208: "Step 5 – Clinical Consultation: Doctor accesses consolidated cross-hospital EHR and records diagnostic notes and prescriptions.",
    209: "Step 6 – Maternal / Child Care Linking: Records are bound to official RCH / PICME references, and pediatric vaccine dates are computed.",
    210: "Step 7 – Inter-Hospital Transfer: Emergency referral summary is dispatched electronically to the receiving hospital's clinical staff.",
    211: "Step 8 – Multi-Channel Reminder Dispatch: Background scheduler triggers automated reminders via Telegram bot, SMS, and email.",
    213: (
        "MongoDB Atlas serves as the document-oriented data persistence layer. Its flexible, schema-governed document model accommodates complex, "
        "nested healthcare records, variable laboratory test payloads, multi-trimester antenatal vitals, and dynamic pediatric immunization schedules "
        "without requiring rigid database schema migrations."
    ),
    216: "Figure 3.4 illustrates the entity-relationship database schema showing collections, primary foreign key references, and embedded clinical subdocuments.",
    218: "The database maintains structured document collections for users, doctors, appointments, slots, maternal records, and child immunizations.",
    223: "Table 3.1 describes the primary MongoDB collections, schema purposes, and key field attributes including contact numbers and cancellation audit details.",
    225: "Algorithm 1 delineates the feature extraction, label encoding, and Gradient Boosting classification steps used to predict appointment urgency and doctor workload.",
    230: "Figure 3.5 demonstrates the machine learning engineering pipeline spanning data ingestion, categorical feature encoding, model evaluation, and Flask deployment.",
    232: (
        "The machine learning engineering pipeline ingests clinical parameters including Patient Age, Gender, Blood Group, Primary Condition, Active "
        "Medications, and Diagnostic Indicators. Categorical variables are preprocessed using serialized scikit-learn LabelEncoders. The dataset is partitioned "
        "into 80% training and 20% validation subsets. Gradient Boosting Classifiers are trained for appointment urgency classification (Emergency, Consultation, "
        "Follow-up, Vaccination, General Checkup) and clinical risk level stratification (Low, Moderate, High, Critical), achieving high diagnostic fidelity."
    ),

    # CHAPTER 4: SYSTEM IMPLEMENTATION
    235: (
        "This chapter documents the concrete engineering realization, component integration, and programmatic execution of the Unified Healthcare platform. "
        "It details the end-to-end implementation across the React 18 presentation layer, Node.js and Express.js REST API services, Python Flask machine "
        "learning microservices, and MongoDB Atlas persistence. Crucial focus is placed on the deployment of interactive appointment cancellation reasoning, "
        "multi-channel (Telegram bot, SMS, and rich HTML email) notification gateways, and emergency inter-hospital referral protocols."
    ),
    237: (
        "The implementation adopts a 3-Tier Model-View-Controller (MVC) architectural design, decoupling client presentation views, Express.js API "
        "business controllers, and MongoDB data models, while bridging computationally intensive predictive tasks to an asynchronous Flask service."
    ),
    239: "Figure 4.1 illustrates the complete Model-View-Controller architecture connecting frontend React views with backend Express controllers and Mongoose data models.",
    241: (
        "The frontend interface incorporates an interactive Appointment Cancellation Modal equipped with quick-select reason chips ('Doctor unavailable', "
        "'Personal emergency', 'Rescheduling needed', 'Symptoms resolved', 'Travel difficulties', 'Other') alongside an optional custom explanation textarea. "
        "Patient profiles allow mobile phone number updates, and doctor/staff dashboards display prominent cancellation reason badges and audit tooltips."
    ),
    244: (
        "The Home / Dashboard provides role-tailored navigation for Patients, Doctors, Hospital Staff, and System Administrators, allowing frictionless "
        "access to ABHA credential verification, appointment booking, medical history inspection, immunization tracking, multi-language switching, and "
        "clinical workload management."
    ),
    245: (
        "The user interface renders high-priority clinical cards, upcoming consultation notices, vaccination milestone due dates, and quick action "
        "buttons, abstracting complex backend REST endpoints into an intuitive, responsive user experience."
    ),
    247: (
        "Upon successful authentication, signed JSON Web Tokens are validated, prompting the React client to render the role-appropriate dashboard layout. "
        "Clinical summary metrics and active patient records are fetched asynchronously from authenticated Express REST endpoints."
    ),
    248: (
        "Patients review active reservations, due vaccines, and past prescriptions; consulting physicians inspect daily appointment schedules and workload "
        "risk scores; clinic administrators manage doctor rosters and overall department operations."
    ),
    249: (
        "The dashboard serves as an operational integration nexus, combining live user session states, real-time appointment queues, and clinical analytics "
        "widgets within a single coherent workspace."
    ),
    253: (
        "Engineered with React 18 functional components and styled with TailwindCSS adhering to an Organic Modernist visual palette (#ba894d warm amber, "
        "#fdfae3 cream). Custom hooks manage asynchronous data fetching, local UI state, and global authentication context."
    ),
    254: "The layout employs responsive CSS flexbox and grid containers that adapt fluidly across desktop, tablet, and smartphone form factors.",
    277: (
        "The Administrator Module enables hospital management to onboard physicians, assign medical departments, configure consultation fees, manage "
        "staff accounts, and generate recurring appointment slot schedules."
    ),
    279: (
        "Administrators log into a dedicated control portal to inspect all registered doctors and departmental rosters. New practitioner accounts are "
        "provisioned, and consultation slot schedules are generated for specified calendar dates and consultation durations."
    ),
    281: (
        "Administrative operations interact with Express.js endpoints (/api/auth, /api/doctors, /api/slots). All payload parameters are strictly validated "
        "against Mongoose schemas prior to executing atomic database updates."
    ),
    293: (
        "The Doctor & Patient Portal coordinates core clinical appointments, digital health identity binding, and regional accessibility. Patients search "
        "for physicians by specialty, view available slots, submit booking requests, authenticate their ABHA ID, and toggle their preferred language. "
        "Doctors review appointment queues, view patient histories across hospitals, and conduct consultations."
    ),
    295: (
        "Patients browse specialist profiles, select convenient time slots, and supply clinical symptoms. The ABHA integration module (AbhaPage.jsx) allows "
        "users to verify and bind their 14-digit ABHA health ID using standard sandbox API specifications. Interface localization powered by react-i18next "
        "provides instantaneous translation across English, Hindi, Tamil, Telugu, Kannada, and Malayalam. Submitted symptom profiles are transmitted to the "
        "Flask ML microservice for urgency triage, allowing physicians to view prioritized patient queues, historical records, and workload risk indices."
    ),
    298: (
        "React UI components render real-time status indicators (Confirmed, Pending, Cancelled, Completed). Doctor dashboards dynamically present calculated "
        "workload risk scores based on cumulative patient clinical complexity."
    ),
    314: (
        "The Unified Electronic Health Record (EHR) module consolidates clinical diagnoses, active prescriptions, laboratory investigation results, "
        "vital signs, and documented drug allergies across all participating hospitals into a single longitudinal record accessible to authorized clinicians."
    ),
    315: (
        "By centralizing disparate medical entries into a unified patient view, the module provides clinicians with comprehensive diagnostic continuity, "
        "ensuring emergency teams have immediate access to critical allergy records and past clinical histories during urgent care transfers."
    ),
    317: (
        "When a patient visits any affiliated healthcare facility, their complete historical record is retrieved using their unique Patient ID. Attending "
        "physicians can inspect past treatments, previous diagnoses, and lab results regardless of which hospital originally logged the encounter."
    ),
    319: (
        "Clinical histories are stored as structured subdocument arrays in MongoDB. The Express backend exposes secure, authenticated endpoints that "
        "aggregate historical records and enforce role-based access control."
    ),
    320: (
        "The unified data repository ensures seamless clinical record continuity during routine doctor consultations, follow-up visits, and urgent "
        "inter-hospital emergency handoffs."
    ),
    331: (
        "The Machine Learning Preprocessing pipeline cleans and transforms raw patient clinical attributes (Age, Gender, Blood Type, Medical Condition, "
        "Active Medications, and Diagnostic Test Results) into structured numerical matrices suitable for classification inference."
    ),
    333: (
        "Categorical text variables are normalized and converted into numerical representations using scikit-learn LabelEncoders. Continuous variables such "
        "as age are standardized before model execution."
    ),
    335: (
        "Implemented using Python pandas and scikit-learn preprocessing utilities. Encoded dictionary artifacts (encoders.pkl, label_encoder.pkl) are "
        "serialized using joblib for low-latency execution within the Flask microservice."
    ),
    336: (
        "Feature selection identified high-relevance clinical variables governing appointment urgency. The resulting feature matrix is supplied to the "
        "classification algorithms during model training and real-time inference."
    ),
    347: (
        "The Urgency Triage Prediction Model automatically classifies incoming consultation requests into one of five urgency categories: Emergency, "
        "Consultation, Follow-up, Vaccination, or General Checkup, based on submitted patient clinical parameters."
    ),
    348: (
        "This automated classification allows clinical staff to identify urgent cases immediately and assists the system in evaluating doctor workload risk "
        "scores across hospital departments."
    ),
    349: (
        "When an appointment is requested, input attributes are dispatched to the Flask /predict-type endpoint. The trained Gradient Boosting Classifier "
        "evaluates the feature vector and returns the predicted urgency category alongside a model confidence score."
    ),
    350: (
        "The prediction sequence processes clinical symptoms using the identical encoding transformations established during training, ensuring data consistency."
    ),
    351: (
        "The encoded feature vector is evaluated by the classification model, and the resulting triage classification is returned in JSON format to the "
        "React frontend for immediate display on patient and doctor dashboards."
    ),
    353: (
        "Multiple classification algorithms were trained and benchmarked, including Gradient Boosting Classifier, Random Forest Classifier, and Decision "
        "Tree Classifier. The Gradient Boosting Classifier demonstrated the highest predictive accuracy and robustness."
    ),
    357: (
        "The classification models were trained on 80% of the patient dataset and validated on the remaining 20% testing holdout. Based on comparative "
        "evaluations of accuracy, precision, and recall, the Gradient Boosting Classifier was selected for production deployment within the Flask microservice."
    ),
    370: (
        "Comprehensive model evaluation assessed classification performance using standard clinical metrics including Accuracy, Precision, Recall, and "
        "F1-Score to ensure dependable triage categorization."
    ),
    372: (
        "The clinical dataset was split 80:20 into training and testing partitions. Model predictions on unseen test samples were systematically compared "
        "against ground-truth clinical classifications."
    ),
    377: (
        "Model performance was verified using classification metrics across all urgency categories, confirming balanced precision and recall for both acute "
        "emergency cases and routine consultation requests."
    ),
    380: (
        "The Gradient Boosting Classifier achieved superior performance across Emergency and Consultation classes compared to baseline decision trees, "
        "ensuring minimal false-negative classifications for critical cases."
    ),
    391: (
        "The Inter-Hospital Referral Engine enables instant, secure transfer of clinical summaries, active medications, blood group, documented allergies, "
        "and diagnostic histories between healthcare facilities during emergency transfers or specialist referrals."
    ),
    393: (
        "The referring physician selects the target hospital and transmits the referral package. The receiving facility immediately receives authorized access "
        "to the patient's complete longitudinal medical summary, eliminating treatment handoff delays."
    ),
    395: (
        "Developed using Express REST endpoints (/api/emergency-transfers, /api/patients). Secure JWT authentication and institutional access controls ensure "
        "confidential patient data is restricted strictly to licensed medical personnel."
    ),
    407: (
        "The Doctor Workload Risk Scoring module calculates a cumulative clinical burden score (0–100) for each doctor based on the volume and severity of "
        "their active appointments, aiding clinic coordinators in schedule balancing."
    ),
    409: (
        "The Flask /doctor-risk-score endpoint processes the patient profiles assigned to a physician's schedule. The model computes individual patient risk "
        "levels (Low, Moderate, High, Critical) and derives an aggregate clinical workload index."
    ),
    411: (
        "The calculated workload score alerts hospital administrators when doctor consultation volumes reach dangerous thresholds, allowing proactive slot "
        "reallocation to prevent physician fatigue and patient wait-time inflation."
    ),
    426: (
        "The Maternal Healthcare module monitors pregnancy progress, antenatal checkups, high-risk flags, and delivery outcomes, anchoring all records to the "
        "mother's official government RCH / PICME registration ID."
    ),
    428: (
        "Maternal records are initialized using the RCH identifier. Attending obstetricians and health workers log antenatal checkup vitals, lab reports, and "
        "fetal growth milestones across any affiliated clinic visited by the mother."
    ),
    430: (
        "Dedicated Express routes (/api/maternal) manage RCH registration, antenatal visit logging, and high-risk case tracking within MongoDB document collections."
    ),
    442: (
        "The Child Immunization Tracker manages pediatric vaccination schedules from birth through early adolescence, automatically generating milestone "
        "schedules and dispatching proactive multi-channel alerts for upcoming doses."
    ),
    444: (
        "When an infant is registered under the mother's RCH profile, a standardized immunization calendar is generated. Healthcare workers log administered "
        "doses with batch numbers and administration timestamps. An automated background scheduler (reminderScheduler.js) powered by node-cron routinely "
        "evaluates upcoming dose deadlines, dispatching instant push notifications via Telegram bot, mobile SMS, and structured HTML emails."
    ),
    446: (
        "Express immunization routes (/api/immunizations) handle schedule queries, status updates (Pending, Administered, Overdue), and automated reminder triggers."
    ),
    468: (
        "The Core REST API layer establishes secure, structured communication channels connecting the React frontend client, Express API application server, "
        "and Python Flask machine learning microservice."
    ),
    470: (
        "Client requests are routed to Express controllers. When machine learning predictions are required, Express dispatches internal HTTP requests to Flask "
        "endpoints (/predict-type, /doctor-risk-score) and serializes responses back to the client in standard JSON format."
    ),
    472: (
        "Express routes implement modular middleware pipelines for JWT token verification, CORS handling, and Mongoose schema validation. Secure error-handling "
        "middleware prevents sensitive system stack traces from leaking to client responses."
    ),
    475: "API Request Flow: User Action → React 18 UI Component → Express REST Controller → JWT / Mongoose Validation → MongoDB / Flask ML Microservice → JSON Response & Telegram/Email Alerts → Client View Update",

    # CHAPTER 5: SYSTEM TESTING
    501: (
        "This chapter outlines the structured quality assurance strategies, verification protocols, and testing methodologies executed to validate the "
        "platform's operational reliability. It covers rigorous unit testing of API controllers, integration testing of inter-service communications, "
        "functional validation of clinical workflows, and accuracy assessments of predictive machine learning models. The documented results demonstrate "
        "that the platform fulfills all predefined clinical, performance, and security criteria."
    ),
    503: (
        "Comprehensive testing verified that the Unified Healthcare Appointment and Immunization Tracking System operates reliably across all integrated "
        "modules. Test suites evaluated user authentication, ABHA ID verification, multi-language UI rendering, slot management, appointment booking, "
        "reasoned cancellations, slot release, longitudinal EHR retrieval, emergency referral data sharing, maternal RCH tracking, pediatric immunization "
        "logging, automated node-cron notifications, Telegram bot push alerts, and ML urgency triage."
    ),
    504: (
        "All test procedures were executed systematically against the live MERN stack application server and Python Flask microservice, confirming expected "
        "functional outputs, database integrity, and correct HTTP response codes."
    ),
    507: "Individual API controllers, Mongoose schema validators, notification utilities, and Flask ML inference functions were tested in isolation using valid and boundary inputs.",
    509: (
        "Inter-service communication was thoroughly tested across the full application stack: React frontend components dispatching booking requests to "
        "Express REST endpoints, validating against MongoDB collections, executing asynchronous Flask ML triage predictions, and triggering simultaneous "
        "Telegram, SMS, and email alerts."
    ),
    511: "End-to-end user journeys—from patient registration and ABHA verification to slot booking, clinical consultation, reasoned cancellation with slot release, and immunization logging—were validated under simulated hospital conditions.",
    513: "Each system feature was systematically verified against its corresponding functional requirement (FR1 through FR15), confirming complete requirement coverage.",
    515: "Express and Flask REST endpoints were tested using Postman and client Axios requests, verifying HTTP status codes, JSON payload schemas, and authentication headers.",
    517: "Machine learning models were evaluated on independent testing holdouts to confirm high classification accuracy, precision, recall, and robust triage categorization.",
    520: "Table 5.1 details the comprehensive validation test cases covering authentication, appointment scheduling, cancellation reasoning, and multi-channel alerts.",

    # CHAPTER 6: CONCLUSION & FUTURE ENHANCEMENTS
    524: (
        "This chapter synthesizes the engineering accomplishments, clinical contributions, and practical outcomes delivered by the Unified Healthcare "
        "Appointment and Immunization Tracking System. It reflects on how the integration of interoperable electronic health records, multi-channel notification "
        "gateways, automated vaccination scheduling, and predictive clinical triage resolves persistent legacy healthcare deficits. Finally, it outlines "
        "strategic trajectories for future technological expansion, including integration with national health registries and IoT telemetry."
    ),
    526: (
        "The Unified Healthcare Appointment and Immunization Tracking System effectively resolves longstanding challenges in traditional healthcare "
        "administration, specifically isolated medical records, inefficient appointment scheduling, untracked childhood vaccinations, linguistic barriers, "
        "and unmonitored physician workloads. By combining a modern MERN stack architecture with national ABHA digital health verification, automated "
        "multi-channel reminders, six-language localization, and a Python Flask machine learning microservice, the system delivers an accountable, "
        "accessible digital health platform."
    ),
    527: (
        "Key operational highlights include digital ABHA credential binding, responsive multi-language UI translation, longitudinal cross-hospital "
        "health record consolidation, real-time emergency referral data transfer, maternal tracking linked to government RCH identifiers, automated pediatric "
        "vaccination calendars, and Gradient Boosting appointment triage. The machine learning service categorizes visit urgency and tracks doctor workload "
        "risk scores, providing clinical leads with actionable data to optimize staff distribution."
    ),
    528: (
        "Constructed using React 18, TailwindCSS, Express.js, Node.js, MongoDB Atlas, node-cron, Python Flask, and scikit-learn, the platform demonstrates "
        "a resilient, modular microservice architecture. Role-based access control and JWT token authentication safeguard sensitive patient data, while "
        "responsive design ensures seamless usability across mobile, tablet, and desktop environments."
    ),
    529: (
        "By uniting portable electronic medical records, predictive clinical intelligence, maternal-child health tracking, and proactive multi-channel "
        "notifications (Telegram bot, SMS, and email), the system establishes a dependable technological foundation for modern digital healthcare delivery."
    ),
    530: (
        "The platform has been rigorously validated across unit, integration, functional, and API testing benchmarks, confirming robust performance, "
        "high availability, and full readiness for real-world clinical deployment."
    ),
    532: "Integration of WebRTC-based secure video telemedicine consultation capabilities for remote patient consultations.",
    533: "Integration of medical IoT device telemetry for continuous remote vital signs monitoring (pulse oximetry, blood pressure, ECG).",
    534: "Deployment of an AI-powered conversational clinical chatbot to assist patients with preliminary symptom self-assessment.",
    535: "Expansion of automated notification gateways to include interactive WhatsApp Business API messaging for vaccination scheduling.",
    536: "Integration of deep learning computer vision models for automated preliminary screening of chest radiographs and CT scans.",
    537: "Implementation of decentralized, blockchain-based immutable audit logging for patient medical record access and consent management.",
    538: "Predictive analytics modules for hospital bed occupancy forecasting, ICU capacity planning, and emergency department resource allocation.",
    539: "Containerization of microservices using Docker and orchestration with Kubernetes for enterprise-scale multi-hospital deployments.",
    540: "Full compliance auditing and certification aligned with global HIPAA, GDPR, and Indian Digital Personal Data Protection (DPDP) standards.",
    541: "Development of dedicated native mobile applications (iOS and Android) tailored for rural Accredited Social Health Activists (ASHA) field workers."
}

updated_count = 0
for idx, new_text in p_map.items():
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        update_para(p, new_text, font_size_pt=11.5)
        updated_count += 1
    else:
        print(f"Index {idx} out of range (max {len(doc.paragraphs)-1})")

print(f"Successfully updated {updated_count} paragraphs in Arial font!")

# Save to target file
doc.save(target_doc_path)
print(f"Successfully saved updated document to: {target_doc_path}")
