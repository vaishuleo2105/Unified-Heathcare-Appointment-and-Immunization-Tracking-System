import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

for p_path in [r'D:\Unified_Healthcare_Project_Report_Updated.docx', r'D:\Unified_Healthcare_Project_Report (1).docx']:
    try:
        doc = docx.Document(p_path)
        print(f"\nDocument: {p_path} (Paragraphs: {len(doc.paragraphs)})")
        for i in range(len(doc.paragraphs) - 15, len(doc.paragraphs)):
            print(f"P{i}: {doc.paragraphs[i].text.strip()[:80]}")
    except Exception as e:
        print(f"Error reading {p_path}: {e}")
