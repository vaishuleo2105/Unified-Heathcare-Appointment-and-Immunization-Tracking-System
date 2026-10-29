import docx

doc_path = r'D:\Unified_Healthcare_Project_Report (1).docx'
doc = docx.Document(doc_path)

print(f"Total tables: {len(doc.tables)}")

for idx, t in enumerate(doc.tables):
    print(f"\n--- Table {idx} (Rows: {len(t.rows)}, Cols: {len(t.columns)}) ---")
    for r_idx, row in enumerate(t.rows[:5]):
        row_text = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"  Row {r_idx}: {row_text[:3]}")
