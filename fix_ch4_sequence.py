import docx
import subprocess

subprocess.run(['taskkill', '/F', '/IM', 'WINWORD.EXE'], capture_output=True)

doc_path = r'D:\Unified_Healthcare_Project_Report.docx'
doc = docx.Document(doc_path)

# Update document body headings to match sequential 4.1 to 4.5
body_mapping = {
    '4.1 USER PORTALS & MULTILINGUAL PRESENTATION LAYER': '4.1 USER PORTALS & MULTILINGUAL PRESENTATION LAYER',
    '4.2 UNIFIED EHR & EMERGENCY INTER-HOSPITAL REFERRAL ENGINE': '4.2 UNIFIED EHR & EMERGENCY INTER-HOSPITAL REFERRAL ENGINE',
    '4.4 MACHINE LEARNING TRIAGE & WORKLOAD RISK MICROSERVICE': '4.3 MACHINE LEARNING TRIAGE & WORKLOAD RISK MICROSERVICE',
    '4.3 MATERNAL & CHILD IMMUNIZATION TRACKING MODULE': '4.4 MATERNAL & CHILD IMMUNIZATION TRACKING MODULE',
    '4.5 CORE REST API & SECURITY INFRASTRUCTURE': '4.5 CORE REST API & SECURITY INFRASTRUCTURE'
}

for p in doc.paragraphs:
    txt = p.text.strip()
    if txt in body_mapping:
        p.text = body_mapping[txt]
        print('Updated Body Heading ->', p.text)

# Update Table 0 (TOC) to match
t0 = doc.tables[0]
toc_entries = [
    ('4.1 User Portals & Multilingual Presentation Layer', '12'),
    ('4.2 Unified EHR & Emergency Inter-Hospital Referral Engine', '13'),
    ('4.3 Machine Learning Triage & Workload Risk Microservice', '14'),
    ('4.4 Maternal & Child Immunization Tracking Module', '15'),
    ('4.5 Core REST API & Security Infrastructure', '16'),
]

row_start_idx = None
for idx, r in enumerate(t0.rows):
    if '4.1' in r.cells[0].text:
        row_start_idx = idx
        break

if row_start_idx is not None:
    for i, (entry_title, entry_page) in enumerate(toc_entries):
        r = t0.rows[row_start_idx + i]
        r.cells[0].text = entry_title
        r.cells[1].text = entry_page
    print('Successfully synchronized Table of Contents!')

doc.save(doc_path)
print('Successfully saved updated D:\\Unified_Healthcare_Project_Report.docx!')
