import docx
import subprocess

# Ensure Word is closed
subprocess.run(['taskkill', '/F', '/IM', 'WINWORD.EXE'], capture_output=True)

doc_path = r'D:\Unified_Healthcare_Project_Report.docx'
doc = docx.Document(doc_path)

# Map sub-headers inside Chapter 4 to clean structured 4.X.Y sub-numberings
sub_heading_map = {
    '4.1 USER PORTALS & MULTILINGUAL PRESENTATION LAYER': '4.1 USER PORTALS & MULTILINGUAL PRESENTATION LAYER',
    'DOCTOR & PATIENT PORTAL, ABHA ID INTEGRATION & MULTILINGUAL SUPPORT': '4.1.2 Digital ABHA ID Integration & Multilingual Support',
    
    '4.2 UNIFIED EHR & EMERGENCY INTER-HOSPITAL REFERRAL ENGINE': '4.2 UNIFIED EHR & EMERGENCY INTER-HOSPITAL REFERRAL ENGINE',
    'INTER-HOSPITAL REFERRAL & EMERGENCY DATA SHARING': '4.2.2 Inter-Hospital Referral & Emergency Data Sharing Engine',
    
    '4.3 MACHINE LEARNING TRIAGE & WORKLOAD RISK MICROSERVICE': '4.3 MACHINE LEARNING TRIAGE & WORKLOAD RISK MICROSERVICE',
    'APPOINTMENT TYPE PREDICTION (ML SERVICE)': '4.3.2 Appointment Urgency Triage Prediction Model',
    'ML MODEL EVALUATION & CLASSIFICATION PERFORMANCE': '4.3.3 ML Model Evaluation & Performance Metrics',
    'DOCTOR WORKLOAD RISK SCORE & SCHEDULE OPTIMIZATION': '4.3.4 Doctor Workload Risk Scoring & Schedule Balancing',
    
    '4.4 MATERNAL & CHILD IMMUNIZATION TRACKING MODULE': '4.4 MATERNAL & CHILD IMMUNIZATION TRACKING MODULE',
    'CHILD IMMUNIZATION TRACKING & AUTOMATED REMINDERS': '4.4.2 Child Immunization Milestone Tracker & Reminders',
    
    '4.5 CORE REST API & SECURITY INFRASTRUCTURE': '4.5 CORE REST API & SECURITY INFRASTRUCTURE'
}

in_ch4 = False
for p in doc.paragraphs:
    txt = p.text.strip()
    if 'SYSTEM IMPLEMENTATION' in txt:
        in_ch4 = True
    elif 'CHAPTER 5' in txt or '5. SYSTEM TESTING' in txt:
        in_ch4 = False
    elif in_ch4:
        txt_upper = txt.upper()
        for key, new_val in sub_heading_map.items():
            if key in txt_upper:
                p.text = new_val
                print(f'Updated -> {new_val}')
                break

# Now update Table 0 (Table of Contents) to reflect the exact 5 main modules
t0 = doc.tables[0]
toc_entries = [
    ('4.1 User Portals & Multilingual Presentation Layer', '12'),
    ('4.2 Unified EHR & Emergency Inter-Hospital Referral Engine', '13'),
    ('4.3 Machine Learning Triage & Workload Risk Microservice', '14'),
    ('4.4 Maternal & Child Immunization Tracking Module', '15'),
    ('4.5 Core REST API & Security Infrastructure', '16'),
]

row_start_idx = None
for idx, r in enumerate(t0.rows):
    if '4.1' in r.cells[0].text:
        row_start_idx = idx
        break

if row_start_idx is not None:
    for i, (entry_title, entry_page) in enumerate(toc_entries):
        r = t0.rows[row_start_idx + i]
        r.cells[0].text = entry_title
        r.cells[1].text = entry_page
    print('Synchronized Table of Contents (Table 0) successfully!')

doc.save(doc_path)
print('Successfully saved updated D:\\Unified_Healthcare_Project_Report.docx!')
