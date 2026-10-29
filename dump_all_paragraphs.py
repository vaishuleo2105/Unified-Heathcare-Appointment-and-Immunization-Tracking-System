import docx
import json

doc_path = r'D:\Unified_Healthcare_Project_Report (1).docx'
doc = docx.Document(doc_path)

output_lines = []
p_data = []

for idx, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    has_drawing = 'w:drawing' in p._element.xml
    style_name = p.style.name if p.style else 'None'
    runs_info = []
    for r in p.runs:
        runs_info.append({
            'text': r.text,
            'font': r.font.name,
            'size': r.font.size.pt if r.font.size else None,
            'bold': r.bold,
            'italic': r.italic
        })
    
    output_lines.append(f"P{idx} [style={style_name}, drawing={has_drawing}]: {text}")
    p_data.append({
        'index': idx,
        'style': style_name,
        'has_drawing': has_drawing,
        'text': text,
        'runs': runs_info
    })

with open('target_doc_dump.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(output_lines))

with open('target_doc_paragraphs.json', 'w', encoding='utf-8') as f:
    json.dump(p_data, f, indent=2, ensure_ascii=False)

print(f"Dumped {len(doc.paragraphs)} paragraphs successfully.")
