import docx

doc = docx.Document(r'D:\FreightOpt_Project_Report.docx')

with open('inspect_details.txt', 'w', encoding='utf-8') as f:
    f.write(f'Total paragraphs: {len(doc.paragraphs)}\n')
    f.write(f'Total tables: {len(doc.tables)}\n\n')
    
    f.write('=== PARAGRAPHS ===\n')
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip():
            runs_info = [(r.text, r.bold, r.italic, r.font.name, r.font.size) for r in p.runs]
            f.write(f'P[{idx}]: text="{p.text}" | align={p.alignment} | style={p.style.name} | runs={runs_info}\n')
            
    f.write('\n=== TABLES ===\n')
    for t_idx, table in enumerate(doc.tables):
        f.write(f'--- Table {t_idx} ---\n')
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, p in enumerate(cell.paragraphs):
                    if p.text.strip():
                        runs_info = [(r.text, r.bold, r.italic, r.font.name, r.font.size) for r in p.runs]
                        f.write(f'T[{t_idx}] R[{r_idx}] C[{c_idx}] P[{p_idx}]: text="{p.text}" | align={p.alignment} | style={p.style.name} | runs={runs_info}\n')
