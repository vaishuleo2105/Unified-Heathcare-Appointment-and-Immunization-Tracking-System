import docx
import re
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc_path = r'D:\Unified_Healthcare_Project_Report (1).docx'
doc = docx.Document(doc_path)

print(f"Loaded {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables.")

def set_run_arial(run, font_size_pt=11.5):
    run.font.name = 'Arial'
    run.font.size = Pt(font_size_pt)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), 'Arial')
    rFonts.set(qn('w:hAnsi'), 'Arial')
    rFonts.set(qn('w:cs'), 'Arial')

# 1. Search and inspect paragraphs mentioning SMS / Twilio / Fast2SMS
p_sms_found = []
for idx, p in enumerate(doc.paragraphs):
    if re.search(r'\b(sms|fast2sms|twilio)\b', p.text, re.IGNORECASE):
        p_sms_found.append((idx, p.text.strip()))

print(f"\n--- Found {len(p_sms_found)} paragraphs mentioning SMS/Twilio/Fast2SMS ---")
for idx, text in p_sms_found:
    print(f"P{idx}: {text[:100]}...")

# 2. Search tables
table_sms_found = []
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            if re.search(r'\b(sms|fast2sms|twilio)\b', cell.text, re.IGNORECASE):
                table_sms_found.append((t_idx, r_idx, c_idx, cell.text.strip()))

print(f"\n--- Found {len(table_sms_found)} table cells mentioning SMS/Twilio/Fast2SMS ---")
for t_idx, r_idx, c_idx, text in table_sms_found:
    print(f"Table {t_idx} [Row {r_idx}, Col {c_idx}]: {text[:80]}...")
