import docx
from docx.shared import Pt
import shutil
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

src_path = r'D:\Unified_Healthcare_Project_Report (1).docx'
test_path = r'D:\Unified_Healthcare_Project_Report_test.docx'

shutil.copyfile(src_path, test_path)
doc = docx.Document(test_path)

p41 = doc.paragraphs[41]
p41.text = ""
run = p41.add_run("Test text in Arial font.")
run.font.name = "Arial"
run.font.size = Pt(11.5)

print("Before save XML:", run._element.xml)

doc.save(test_path)

doc2 = docx.Document(test_path)
run2 = doc2.paragraphs[41].runs[0]
print("After save XML:", run2._element.xml)
print("run2 font name:", run2.font.name)

if os.path.exists(test_path):
    os.remove(test_path)
print("Done!")
