import docx

doc = docx.Document(r'D:\Unified_Healthcare_Project_Report.docx')

print("Total paragraphs:", len(doc.paragraphs))
print("Total tables:", len(doc.tables))

# Check for drawings
drawings = 0
for p in doc.paragraphs:
    if 'w:drawing' in p._element.xml:
        drawings += 1
print(f"Drawings count: {drawings} / 25")

# Check student names
print("\nChecking student names in document:")
found_names = False
for idx, p in enumerate(doc.paragraphs):
    if "INDUJA" in p.text or "VAISHALI" in p.text or "Induja" in p.text or "Vaishali" in p.text:
        print(f"P[{idx}]: {p.text}")
        found_names = True

# Check for remaining Freight / FreightOpt references
print("\nChecking for any remaining FreightOpt references:")
freight_count = 0
for idx, p in enumerate(doc.paragraphs):
    if "Freight" in p.text or "freight" in p.text:
        print(f"P[{idx}]: {p.text[:100]}...")
        freight_count += 1

for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            if "Freight" in cell.text or "freight" in cell.text:
                print(f"Table[{t_idx}] R[{r_idx}] C[{c_idx}]: {cell.text[:100]}...")
                freight_count += 1

print(f"Remaining FreightOpt occurrences: {freight_count}")
