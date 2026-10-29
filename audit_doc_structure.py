import docx

doc = docx.Document(r'D:\Unified_Healthcare_Project_Report.docx')

with open('doc_structure.txt', 'w', encoding='utf-8') as f:
    f.write(f'Total Paragraphs: {len(doc.paragraphs)}\n')
    f.write(f'Total Tables: {len(doc.tables)}\n\n')
    for i, p in enumerate(doc.paragraphs):
        f.write(f'P[{i}]: text="{p.text[:120]}" | style={p.style.name} | align={p.alignment}\n')

print(f'Done! Evaluated {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables.')
