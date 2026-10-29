import docx
import os

target_path = r'D:\Unified_Healthcare_Project_Report (1).docx'
print(f"Checking target file: {target_path}")
print(f"Exists: {os.path.exists(target_path)}")

if os.path.exists(target_path):
    doc = docx.Document(target_path)
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total tables: {len(doc.tables)}")
    
    # Check fonts of existing paragraphs
    fonts = set()
    for p in doc.paragraphs[:50]:
        for r in p.runs:
            if r.font.name:
                fonts.add(r.font.name)
    print(f"Sample fonts detected in runs: {fonts}")
    
    # Check default style font
    try:
        print(f"Normal style font: {doc.styles['Normal'].font.name}")
    except Exception as e:
        print(f"Normal style font query error: {e}")
