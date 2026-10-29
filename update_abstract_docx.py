import docx

doc_path = 'Project_Abstract_Template1.docx'
doc = docx.Document(doc_path)

# Proposed Work text updated with ML microservice, emergency transfer, audit logs, and ABHA integration details
proposed_work_text = (
    "               This project proposes a web-based Healthcare Appointment and Immunization Tracking System "
    "built on the MERN stack (MongoDB, Express.js, React.js, Node.js) styled with TailwindCSS following Organic Modernist "
    "design principles, with integrated Python ML risk scoring, ABHA digital health ID verification, and multi-language support "
    "(English, Hindi, Tamil, Telugu, Kannada, Malayalam). The system enables patients to search for and book appointments "
    "with healthcare providers while sending automated email reminders for check-ups and immunization schedules. At its core, "
    "the platform maintains a unified electronic health record (EHR) consolidating diagnoses, treatments, prescriptions, and "
    "vaccination records across hospitals. For critical referrals, an Emergency Transfer module dispatches patient medical snapshots "
    "to receiving hospitals alongside automated security audit logging. For maternal and child care, the system links antenatal "
    "checkups, delivery details, and child immunization schedules under government-issued reference numbers."
)

deliverables_text = (
    "              A fully functional web application comprising: patient registration & ABHA authentication portal; "
    "appointment booking and time-slot management module; automated email reminder service; unified electronic health record "
    "(EHR) module with downloadable PDF summaries; inter-hospital emergency referral dispatch module; security audit logging; "
    "maternal and child immunization tracking module; multi-role clinical dashboards (Patient, Doctor, Staff, Admin); "
    "and a comprehensive project report documenting system architecture, database schema, and implementation details."
)

# Update paragraphs 8 and 12 in docx
if len(doc.paragraphs) > 8:
    doc.paragraphs[8].text = proposed_work_text
if len(doc.paragraphs) > 12:
    doc.paragraphs[12].text = deliverables_text

doc.save('Project_Abstract_Template1.docx')
print('Successfully updated Project_Abstract_Template1.docx')
